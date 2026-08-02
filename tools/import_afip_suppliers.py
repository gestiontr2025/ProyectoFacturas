"""Generar el catálogo JSON desde un Excel de comprobantes recibidos.

Uso:
    python tools/import_afip_suppliers.py "ruta/al/archivo.xlsx"

El Excel es una fuente de importación. El programa principal nunca necesita
abrirlo: durante la ejecución diaria utiliza exclusivamente el JSON generado.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date
import json
from pathlib import Path
import re
import sys
import unicodedata

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = PROJECT_ROOT / "suppliers" / "data" / "supplier_catalog.json"
RECEIVER_CUIT = "30-71834746-3"


def normalize_text(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(c for c in text if unicodedata.category(c) != "Mn").upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def normalize_cuit(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = re.sub(r"\D", "", str(value))
    if len(digits) != 11:
        return None
    return f"{digits[:2]}-{digits[2:10]}-{digits[10]}"


def slugify(value: str) -> str:
    return re.sub(r"_+", "_", normalize_text(value).lower().replace(" ", "_")).strip("_")


def read_existing_catalog(path: Path) -> dict[str, dict]:
    """Indexar el catálogo anterior para conservar alias y nombres de carpeta."""
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    index: dict[str, dict] = {}
    for item in data.get("suppliers", []):
        for candidate in [item.get("legal_name"), item.get("display_name"), *item.get("aliases", [])]:
            if candidate:
                index.setdefault(normalize_text(candidate).replace(" ", ""), item)
    return index


def detect_header_row(sheet) -> tuple[int, dict[str, int]]:
    required = {"Nro. Doc. Emisor", "Denominación Emisor", "Tipo", "Moneda"}
    for row_number in range(1, min(sheet.max_row, 25) + 1):
        values = [cell.value for cell in sheet[row_number]]
        mapping = {str(value).strip(): index for index, value in enumerate(values) if value is not None}
        if required.issubset(mapping):
            return row_number, mapping
    raise ValueError("No se encontró la fila de encabezados esperada en el Excel.")


def import_catalog(input_path: Path, output_path: Path, minimum_occurrences: int) -> dict:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook.active
    header_row, columns = detect_header_row(sheet)

    groups: dict[str, dict] = {}
    source_rows = 0
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        source_rows += 1
        cuit = normalize_cuit(values[columns["Nro. Doc. Emisor"]])
        legal_name = values[columns["Denominación Emisor"]]
        if not cuit or not legal_name or cuit == RECEIVER_CUIT:
            continue
        key = cuit.replace("-", "")
        group = groups.setdefault(key, {
            "names": Counter(), "types": set(), "currencies": set(),
            "vat_rates": set(), "other_taxes": False, "count": 0,
        })
        group["count"] += 1
        group["names"][str(legal_name).strip()] += 1
        document_type = values[columns["Tipo"]]
        if document_type:
            group["types"].add(str(document_type).strip())
        currency = values[columns["Moneda"]]
        if currency:
            group["currencies"].add({"$": "ARS", "U$S": "USD"}.get(str(currency).strip(), str(currency).strip()))
        for column_name, rate in (("IVA 2,5%", 2.5), ("IVA 5%", 5.0), ("IVA 10,5%", 10.5), ("IVA 21%", 21.0), ("IVA 27%", 27.0)):
            if column_name in columns:
                value = values[columns[column_name]]
                if isinstance(value, (int, float)) and value != 0:
                    group["vat_rates"].add(rate)
        if "Otros Tributos" in columns:
            value = values[columns["Otros Tributos"]]
            group["other_taxes"] = group["other_taxes"] or (isinstance(value, (int, float)) and value != 0)

    previous = read_existing_catalog(output_path)
    suppliers = []
    used_ids: set[str] = set()
    for digits, group in sorted(groups.items(), key=lambda item: (-item[1]["count"], item[1]["names"].most_common(1)[0][0])):
        if group["count"] < minimum_occurrences:
            continue
        legal_name = group["names"].most_common(1)[0][0]
        old = previous.get(normalize_text(legal_name).replace(" ", ""), {})
        identifier = old.get("identifier") or slugify(legal_name)
        if identifier in used_ids:
            identifier = f"{identifier}_{digits[-4:]}"
        used_ids.add(identifier)
        aliases = set(old.get("aliases", []))
        aliases.update(name for name in group["names"] if normalize_text(name) != normalize_text(legal_name))
        suppliers.append({
            "identifier": identifier,
            "legal_name": legal_name,
            "display_name": old.get("display_name"),
            "cuit": f"{digits[:2]}-{digits[2:10]}-{digits[10]}",
            "folder_name": old.get("folder_name"),
            "aliases": sorted(aliases, key=normalize_text),
            "observed": {
                "document_types": sorted(group["types"]),
                "vat_rates": sorted(group["vat_rates"]),
                "other_taxes": bool(group["other_taxes"]),
                "currencies": sorted(group["currencies"]),
                "document_count": group["count"],
            },
            "source": "AFIP/ARCA - Mis Comprobantes Recibidos",
        })

    result = {
        "schema_version": 1,
        "generated_at": date.today().isoformat(),
        "source": {
            "type": "AFIP_ARCA_RECEIVED_DOCUMENTS_XLSX",
            "receiver_cuit": RECEIVER_CUIT,
            "minimum_occurrences": minimum_occurrences,
            "source_rows": source_rows,
            "unique_emitters": len(groups),
            "included_suppliers": len(suppliers),
            "excluded_single_occurrence_emitters": sum(1 for group in groups.values() if group["count"] < minimum_occurrences),
        },
        "suppliers": suppliers,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(output_path)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Importar proveedores recurrentes desde un Excel de AFIP/ARCA.")
    parser.add_argument("excel", type=Path, help="Ruta del Excel Mis Comprobantes Recibidos.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Ruta del catálogo JSON generado.")
    parser.add_argument("--min-occurrences", type=int, default=2, help="Cantidad mínima de comprobantes por emisor.")
    args = parser.parse_args()
    if args.min_occurrences < 1:
        parser.error("--min-occurrences debe ser al menos 1.")
    if not args.excel.is_file():
        parser.error(f"No se encontró el Excel: {args.excel}")
    try:
        result = import_catalog(args.excel, args.output, args.min_occurrences)
    except Exception as exc:
        print(f"No fue posible importar el catálogo: {exc}", file=sys.stderr)
        return 1
    source = result["source"]
    print("Catálogo generado correctamente.")
    print(f"Proveedores incluidos: {source['included_suppliers']}")
    print(f"Emisores excluidos por baja frecuencia: {source['excluded_single_occurrence_emitters']}")
    print(f"Archivo: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
