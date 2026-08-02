"""Generación del Excel de perfil impositivo por proveedor."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from taxes.profile import build_supplier_tax_profiles


HEADERS = (
    "Nombre Proveedor",
    "CUIT Proveedor",
    "IVA 27%",
    "IVA 21%",
    "IVA 10,5%",
    "Per. IVA",
    "Per. IIBB CABA",
    "Per. IIBB BSAS",
    "Impuestos internos",
)


@dataclass(frozen=True)
class ExportResult:
    output_path: Path
    suppliers: int
    invoices_analyzed: int
    invoices_skipped: int


def _desktop_folder() -> Path:
    """Localizar el Escritorio en instalaciones normales o con OneDrive."""
    candidates = []
    onedrive = os.getenv("OneDrive") or os.getenv("OneDriveCommercial")
    if onedrive:
        candidates.append(Path(onedrive) / "Desktop")
        candidates.append(Path(onedrive) / "Escritorio")
    candidates.extend((Path.home() / "Desktop", Path.home() / "Escritorio"))
    for candidate in candidates:
        if candidate.exists():
            return candidate
    fallback = Path.home() / "Desktop"
    fallback.mkdir(parents=True, exist_ok=True)
    return fallback


def _yes_no(value: bool) -> str:
    return "Sí" if value else "No"


def export_supplier_tax_profile(invoice_root: Path, output_path: Path | None = None) -> ExportResult:
    profiles, analyzed, skipped = build_supplier_tax_profiles(Path(invoice_root))
    if output_path is None:
        timestamp = datetime.now().strftime("%Y-%m-%d")
        output_path = _desktop_folder() / f"Perfil_Impositivo_Proveedores_{timestamp}.xlsx"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Perfil impositivo"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    sheet.append(HEADERS)
    for profile in profiles:
        sheet.append(
            (
                profile.supplier_name,
                profile.supplier_cuit,
                _yes_no(profile.vat_27),
                _yes_no(profile.vat_21),
                _yes_no(profile.vat_10_5),
                _yes_no(profile.vat_perception),
                _yes_no(profile.iibb_caba_perception),
                _yes_no(profile.iibb_bsas_perception),
                _yes_no(profile.internal_taxes),
            )
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    yes_fill = PatternFill("solid", fgColor="E2F0D9")
    no_fill = PatternFill("solid", fgColor="F2F2F2")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    sheet.row_dimensions[1].height = 32

    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=9):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
            cell.fill = yes_fill if cell.value == "Sí" else no_fill

    widths = {
        "A": 42,
        "B": 18,
        "C": 12,
        "D": 12,
        "E": 13,
        "F": 12,
        "G": 17,
        "H": 17,
        "I": 20,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    if profiles:
        table = Table(displayName="SupplierTaxProfileTable", ref=f"A1:I{len(profiles) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    workbook.save(output_path)
    return ExportResult(
        output_path=output_path,
        suppliers=len(profiles),
        invoices_analyzed=analyzed,
        invoices_skipped=skipped,
    )
